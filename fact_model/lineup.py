"""Line-up facts: SNAPSHOT math only.

Each report_date is an independent snapshot of tonnage queued. Snapshots are plotted
as-is over their real dates; they are never accumulated or differenced. Cumulative
(executed) math lives in executed.py.
"""

import pandas as pd

from .dimensions import add_geography, commodity_group, default_config, marketing_year_label

DIMS = ["report_date", "source", "region", "port", "elevator", "commodity", "destination"]
FACT_COLS = DIMS + ["month", "marketing_year", "kmt", "vessels", "vessels_no_tonnage",
                    "mapped", "unmapped_reason"]


def _status_norm(s):
    # Mirrors app.load_data so the new fact agrees with the existing tabs.
    if pd.isna(s):
        return "Unknown"
    s = str(s).strip().upper()
    if s.startswith("SAIL"):
        return "Sailed"
    if s in ("LOADING", "PART-LDD", "IN PORT"):
        return "Loading/In Port"
    if s == "FILED":
        return "Filed"
    if s.startswith("ETA") or s.startswith("L/R"):
        return "ETA"
    return "Other"


_US_COLS = ["ELEVATOR", "VESSEL", "ATA", "STATUS", "MT", "COMMODITY", "DESTINATION", "SAIL DATE"]


def _vessels_from_sheet(sheet, d):
    d = d.dropna(how="all")
    return pd.DataFrame({
        "sheet": sheet,
        "elevator_raw": d["ELEVATOR"],
        "vessel": d["VESSEL"],
        "status_norm": d["STATUS"].apply(_status_norm),
        "kmt": pd.to_numeric(d["MT"], errors="coerce"),
        "commodity_raw": d["COMMODITY"],
        "destination": d["DESTINATION"],
        "sail_date": pd.to_datetime(d["SAIL DATE"], errors="coerce"),
    })


def read_us_vessels(path):
    """Vessel-level US frame from the dashboard workbook (USG/PNW/TXG sheets)."""
    frames = []
    for sheet in ("USG", "PNW", "TXG"):
        d = pd.read_excel(path, sheet_name=sheet, usecols=range(8), header=0)
        d.columns = _US_COLS
        frames.append(_vessels_from_sheet(sheet, d))
    return pd.concat(frames, ignore_index=True)


def read_southport_vessels(path):
    """Vessel-level frame straight from a raw Southport email workbook (MISS RIVER / TEXAS / PNW
    SORT + SAIL sheets), using convert_southport.read_region so historical files are read exactly
    as the live converter reads them."""
    import importlib
    from openpyxl import load_workbook
    cs = importlib.import_module("convert_southport")
    pairs = {"USG": ("MISS RIVER SORT", "MISS RIVER SAIL", False),
             "TXG": ("TEXAS SORT", "TEXAS SAIL", True),
             "PNW": ("PNW SORT", "PNW SAIL", False)}
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        frames = [_vessels_from_sheet(sheet, cs.read_region(wb, a, b, texas=t))
                  for sheet, (a, b, t) in pairs.items()]
    finally:
        wb.close()
    return pd.concat(frames, ignore_index=True)


def _aggregate(v, report_date_col="report_date"):
    v = v.copy()
    v["destination"] = (v["destination"].astype("string").str.strip().str.upper()
                        .replace({"RVT": "UNKNOWN", "?": "UNKNOWN", "": "UNKNOWN"})
                        .fillna("UNKNOWN"))
    v["no_tonnage"] = v["kmt"].isna()
    v["kmt"] = v["kmt"].fillna(0.0)
    v["vessel_id"] = v["port"].astype(str) + "|" + v["elevator"].astype(str) + "|" + v["vessel"].astype(str)
    g = (v.groupby(DIMS + ["mapped", "unmapped_reason"], dropna=False)
          .agg(kmt=("kmt", "sum"), vessels=("vessel_id", "nunique"),
               vessels_no_tonnage=("no_tonnage", "sum")).reset_index())
    g["report_date"] = pd.to_datetime(g["report_date"])
    g["month"] = g["report_date"].dt.to_period("M").dt.to_timestamp()
    g["marketing_year"] = [marketing_year_label(d, c) for d, c in zip(g["report_date"], g["commodity"])]
    return g[FACT_COLS].sort_values(DIMS).reset_index(drop=True)


def us_lineup_snapshot(vessels, report_date, cfg=None):
    """Snapshot of US tonnage queued as of report_date. Queued = not sailed, no sail date,
    and a real vessel record. The TXG sheet is padded with dock-name-only rows (no vessel,
    status or tonnage); those are placeholders, not queue, and are excluded."""
    cfg = cfg or default_config()
    q = vessels[(vessels["status_norm"] != "Sailed") & vessels["sail_date"].isna()
                & (vessels["vessel"].notna() | (vessels["status_norm"] != "Unknown"))].copy()
    q = add_geography(q, "US", cfg)
    q["source"] = "US"
    q["report_date"] = pd.Timestamp(report_date)
    q["commodity"] = [commodity_group("US", c) for c in q["commodity_raw"]]
    return _aggregate(q)


def brazil_lineup_snapshot(lineup_df, cfg=None):
    """lineup_df: rows from brazil_db.get_all_lineup (all snapshots). One fact snapshot per
    report_date. Excluded rows (MAINTENANCE etc.) are dropped by the parser's own flag."""
    cfg = cfg or default_config()
    q = lineup_df[lineup_df["excluded"] == 0].copy()
    q = q.rename(columns={"port": "port_raw"})
    q = add_geography(q, "Brazil", cfg)
    q["source"] = "Brazil"
    q["kmt"] = pd.to_numeric(q["mt"], errors="coerce") / 1000.0
    q["commodity"] = [commodity_group("Brazil", p) for p in q["product"]]
    return _aggregate(q)


def asof_snapshot(fact, as_of=None):
    """The single most recent snapshot per source at or before as_of."""
    d = fact if as_of is None else fact[fact["report_date"] <= pd.Timestamp(as_of)]
    last = d.groupby("source")["report_date"].transform("max")
    return d[d["report_date"] == last]


def lineup_trend(fact, by=("commodity",)):
    """Queued tonnage per real report_date, split by `by`. No resampling, no accumulation:
    the x-axis is exactly the snapshot dates we hold (uneven spacing preserved)."""
    by = list(by)
    return (fact.groupby(["report_date", *by], dropna=False)[["kmt", "vessels"]]
                .sum().reset_index().sort_values(["report_date", *by]))
