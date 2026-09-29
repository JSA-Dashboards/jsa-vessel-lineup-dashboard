"""
convert_southport.py
Converts a Southport vessel-lineup Excel into the format the dashboard expects,
then commits and pushes to GitHub.

Southport layout            → Dashboard layout
────────────────────────────────────────────────
MISS RIVER SORT + SAIL      → USG sheet
TEXAS SORT + TEXAS SAIL     → TXG sheet
PNW SORT   + PNW SAIL       → PNW sheet
Trends sheet preserved from existing Vessel Lineup - US.xlsx
"""

import os
import sys
import subprocess
import shutil

import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils.dataframe import dataframe_to_rows

REPO_DIR    = os.path.dirname(os.path.abspath(__file__))
DEST_FILE   = os.path.join(REPO_DIR, "Vessel Lineup - US.xlsx")
BACKUP_FILE = os.path.join(REPO_DIR, "Vessel Lineup - US.bak.xlsx")

# Expected output columns for every US region sheet
OUT_COLS = ["ELEVATOR", "VESSEL", "ATA", "STATUS", "MT", "COMMODITY", "DESTINATION", "SAIL DATE"]


def _clean_mt(val):
    """Convert '35K' → 35, keep 'RVT' as-is, pass numeric through."""
    if pd.isna(val):
        return val
    s = str(val).strip().upper()
    if s in ("RVT", ""):
        return s
    if s.endswith("K"):
        try:
            return float(s[:-1])
        except ValueError:
            pass
    try:
        return float(s)
    except ValueError:
        return val


def read_region(sp_wb, sort_sheet, sail_sheet, texas=False):
    """
    Read a pair of SORT + SAIL sheets from the Southport workbook and
    return a single DataFrame with OUT_COLS.
    """
    def _sheet_to_df(sheet_name):
        ws = sp_wb[sheet_name]
        data = [row for row in ws.iter_rows(values_only=True)
                if any(c is not None for c in row)]
        if len(data) < 2:
            return pd.DataFrame(columns=OUT_COLS)

        header = [str(c).strip().upper() if c is not None else "" for c in data[0]]
        rows   = data[1:]

        if texas:
            # Two ELEVATOR columns — merge as "TERMINAL / ELEVATOR"
            # Col 0 = terminal/dock, Col 1 = grain company
            df = pd.DataFrame(rows, columns=header[:len(rows[0])] if rows else header)
            elev_cols = [c for c in df.columns if c == "ELEVATOR"]
            if len(elev_cols) >= 2:
                cols = list(df.columns)
                first_idx  = cols.index("ELEVATOR")
                # rename second occurrence
                cols[first_idx + 1 + cols[first_idx+1:].index("ELEVATOR")] = "ELEVATOR2"
                df.columns = cols
                df["ELEVATOR"] = df.apply(
                    lambda r: (
                        f"{r['ELEVATOR']} / {r['ELEVATOR2']}"
                        if pd.notna(r.get("ELEVATOR2")) and str(r.get("ELEVATOR2","")).strip()
                        else str(r["ELEVATOR"])
                    ),
                    axis=1,
                )
                df = df.drop(columns=["ELEVATOR2"], errors="ignore")
        else:
            df = pd.DataFrame(rows, columns=header[:len(rows[0])] if rows else header)

        # PNW SAIL uses "SA" instead of "STATUS"
        if "SA" in df.columns and "STATUS" not in df.columns:
            df = df.rename(columns={"SA": "STATUS"})

        return df

    sort_df = _sheet_to_df(sort_sheet)
    sail_df = _sheet_to_df(sail_sheet)

    # SORT rows have no SAIL DATE — add blank column so concat works cleanly
    if "SAIL DATE" not in sort_df.columns:
        sort_df["SAIL DATE"] = None

    combined = pd.concat([sort_df, sail_df], ignore_index=True)

    # Keep only the expected columns; fill any missing ones with None
    for col in OUT_COLS:
        if col not in combined.columns:
            combined[col] = None
    combined = combined[OUT_COLS]

    # Normalise MT values
    combined["MT"] = combined["MT"].apply(_clean_mt)

    # Drop completely empty rows
    combined = combined.dropna(how="all")

    return combined


def git_push(message):
    cmds = [
        ["git", "-C", REPO_DIR, "add",    "Vessel Lineup - US.xlsx"],
        ["git", "-C", REPO_DIR, "commit", "-m", message,
         "--author", "Kolten Postin <275148418+koltenpostin93-blip@users.noreply.github.com>"],
        ["git", "-C", REPO_DIR, "push"],
    ]
    for cmd in cmds:
        r = subprocess.run(cmd, capture_output=True, text=True)
        if r.returncode != 0:
            if "nothing to commit" in r.stdout + r.stderr:
                print("  (no changes to commit)")
                return
            print(f"ERROR: {' '.join(cmd)}")
            print(r.stderr or r.stdout)
            input("\nPress Enter to exit.")
            sys.exit(1)


def convert(southport_path, push=True):
    print(f"\nConverting: {os.path.basename(southport_path)}")

    sp_wb = load_workbook(southport_path, read_only=True, data_only=True)

    regions = {
        "USG": ("MISS RIVER SORT", "MISS RIVER SAIL", False),
        "TXG": ("TEXAS SORT",      "TEXAS SAIL",      True),
        "PNW": ("PNW SORT",        "PNW SAIL",        False),
    }

    dfs = {}
    for region, (sort_sh, sail_sh, is_texas) in regions.items():
        df = read_region(sp_wb, sort_sh, sail_sh, texas=is_texas)
        print(f"  {region}: {len(df)} rows  "
              f"({len(df[df['SAIL DATE'].isna()])} lined up, "
              f"{len(df[df['SAIL DATE'].notna()])} sailed)")
        dfs[region] = df

    sp_wb.close()

    # Read Trends sheet from the existing file (Southport doesn't have it)
    trends_df = None
    if os.path.exists(DEST_FILE):
        try:
            existing_wb = load_workbook(DEST_FILE, read_only=True, data_only=True)
            if "Trends" in existing_wb.sheetnames:
                ws_t = existing_wb["Trends"]
                trends_rows = list(ws_t.iter_rows(values_only=True))
                trends_df = pd.DataFrame(trends_rows)
                print(f"  Trends: {len(trends_df)} rows preserved from existing file")
            existing_wb.close()
        except Exception as e:
            print(f"  WARNING: Could not read Trends from existing file: {e}")

    # Back up the existing file
    if os.path.exists(DEST_FILE):
        shutil.copy2(DEST_FILE, BACKUP_FILE)

    # Build the new workbook using openpyxl
    from openpyxl import Workbook
    new_wb = Workbook()
    new_wb.remove(new_wb.active)  # remove default sheet

    for region, df in dfs.items():
        ws = new_wb.create_sheet(region)
        for r in dataframe_to_rows(df, index=False, header=True):
            ws.append(r)

    if trends_df is not None:
        ws_t = new_wb.create_sheet("Trends")
        for row in trends_df.itertuples(index=False, name=None):
            ws_t.append(list(row))

    new_wb.save(DEST_FILE)
    print(f"\nSaved: {DEST_FILE}")

    if push:
        src_name = os.path.basename(southport_path)
        print("Pushing to GitHub...")
        git_push(f"Update vessel lineup from Southport: {src_name}")
        print("Done! Streamlit Cloud will refresh in ~1 minute.")

    return True


def main():
    # Allow path as CLI arg, otherwise look in Downloads for the latest Southport file
    if len(sys.argv) > 1:
        southport_path = sys.argv[1]
    else:
        downloads = os.path.join(os.path.expanduser("~"), "Downloads")
        candidates = [
            f for f in os.listdir(downloads)
            if "southport" in f.lower() and f.lower().endswith(".xlsx")
        ]
        if not candidates:
            print("ERROR: No Southport Excel file found in Downloads.")
            print("Save the email attachment to Downloads, then run this script again.")
            input("\nPress Enter to exit.")
            sys.exit(1)
        # Pick the most recently modified
        candidates.sort(
            key=lambda f: os.path.getmtime(os.path.join(downloads, f)),
            reverse=True,
        )
        southport_path = os.path.join(downloads, candidates[0])
        print(f"Found: {candidates[0]}")

    convert(southport_path, push=True)
    input("\nPress Enter to exit.")


if __name__ == "__main__":
    main()
