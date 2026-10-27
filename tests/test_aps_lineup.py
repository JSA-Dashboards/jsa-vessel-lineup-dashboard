"""APS 'Grain Elevator Line-ups' history -> US Gulf / PNW snapshots on the Southport basis.

    pytest tests/test_aps_lineup.py -v
"""

import os
import sys

import numpy as np
import pandas as pd
import pytest
from openpyxl import Workbook

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from fact_model import aps_lineup as ap
from fact_model import changes as ch
from fact_model import store


def test_cargo_parsing_covers_the_three_aps_styles():
    f = ap.parse_cargo
    assert f("YC SBM", "30K") == (30.0, "CORN/SBM", False)
    assert f("66K WHEAT") == (66.0, "WHT", False)
    assert f("62K YC/WHEAT") == (62.0, "CORN/WHT", False)
    assert f("65,000 MT CORN") == (65.0, "CORN", False)
    k, c, rvt = f("YC YSB SRW NSW SBM", "10K")
    assert (k, c, rvt) == (10.0, "CORN/YSB/WHT/SBM", False)
    k, c, rvt = f("REVERTING")
    assert np.isnan(k) and c == "RVT" and rvt
    assert f("38K WHT/BY-PROD") == (38.0, "WHT/BY-PROD", False)


def workbook(path, date="2025-09-12"):
    wb = Workbook()
    ws = wb.active
    ws.title = "Miss River"
    ws.append([None, "CLEAR"] + [None] * 7); ws[1][0].value = pd.Timestamp(date).to_pydatetime()
    ws.append([None, "Vessel Name", "ETA", "ETB", "ETC", "Status", "Cargo", "Tons", "Destination"])
    ws.append(["ADM Ama", "STAR OPAL", "IN PORT", "IN BERTH", None, "LOADING", "YC", "50K", "JAPAN"])
    ws.append([None, "FOUR TURANDOT", None, None, None, None, "YC SBM", "30K", "MEXICO"])
    ws.append(["CHS", "ELISA", None, None, None, None, "REVERTING", None, "REVERTING"])
    ws.append(["ADM-BUOYS", "NONE", None, None, None, None, None, None, None])
    ws2 = wb.create_sheet("Puget Sound")
    ws2.append(["Elevator", "Vessel Name", "ETA", "Status", "ETB", "ETD", "Cargo", "Destination"])
    ws2.append(["LOUIS DRYFUS", "IDLE"])
    ws2.append(["TEMCO", "RYSY", None, None, None, None, "70,000 MT CORN", "KOREA"])
    wb.save(path)


def test_workbook_becomes_southport_shaped_vessel_rows_and_snapshots(tmp_path):
    p = str(tmp_path / "2025-09-12__lineups.xlsx")
    workbook(p)
    rd, v = ap.read_workbook(p, "2025-09-12")
    assert rd == pd.Timestamp("2025-09-12")
    assert v["vessel"].tolist() == ["STAR OPAL", "FOUR TURANDOT", "ELISA", "RYSY"]                      # IDLE / NONE rows dropped
    assert v.loc[v["vessel"] == "ELISA", "commodity_raw"].iloc[0] == "RVT" and v["kmt"].notna().sum() == 3
    assert set(v["sheet"]) == {"USG", "PNW"} and v["elevator_raw"].iloc[1] == "ADM Ama"             # elevator carried down its block
    _, snap = ap.snapshot(p, "2025-09-12")
    assert snap["mapped"].sum() >= 1 and set(snap["region"]) == {"US Gulf", "PNW"}
    assert snap.loc[(snap["commodity"] == "Mixed Cargo"), "combo"].iloc[0] == "CORN/SBM"
    assert snap["vessels"].sum() == pytest.approx(4.0)


def snap(date, region, v, k):
    return pd.DataFrame([{"report_date": pd.Timestamp(date), "source": "US", "region": region, "port": "P", "elevator": "E", "commodity": "Corn",
                          "combo": "", "destination": "X", "month": pd.Timestamp(date).replace(day=1), "marketing_year": "2025/26", "kmt": k,
                          "vessels": v, "vessels_no_tonnage": 0.0, "mapped": True, "unmapped_reason": ""}])


def test_overlap_scales_and_scaling_put_aps_on_the_southport_basis():
    aps = {pd.Timestamp("2026-01-10"): pd.concat([snap("2026-01-10", "PNW", 84, 920), snap("2026-01-10", "US Gulf", 103, 1050)])}
    facts = pd.concat([snap("2026-01-10", "PNW", 100, 1000), snap("2026-01-10", "US Gulf", 100, 1000)])
    sc = ap.overlap_scales(aps, facts)
    assert sc["PNW"][:2] == (pytest.approx(0.84), pytest.approx(0.92)) and sc["US Gulf"][0] == pytest.approx(1.03)
    out = ap.scale_snapshot(aps[pd.Timestamp("2026-01-10")], sc)
    assert out.loc[out["region"] == "PNW", "vessels"].iloc[0] == pytest.approx(100) and out.loc[out["region"] == "PNW", "kmt"].iloc[0] == pytest.approx(1000)


def test_backfill_loads_only_uncovered_dates_scaled_and_is_idempotent(tmp_path):
    folder = tmp_path / "aps"
    folder.mkdir()
    for d in ("2024-05-01", "2024-05-02", "2026-01-10"):
        workbook(str(folder / f"{d}__lineups.xlsx"), d)
    db = str(tmp_path / "f.db")
    store.save_lineup_snapshot(pd.concat([snap("2026-01-10", "PNW", 1, 70), snap("2026-01-10", "US Gulf", 3, 100)]), db)       # Southport day
    r = ap.backfill(str(folder), db, since="2024-01-01")
    assert r["to_load"] == 2 and r["first"] == "2024-05-01" and r["scales"]["PNW"][2] == 1
    lu = store.load_lineup(db)
    assert sorted(lu["report_date"].dt.strftime("%Y-%m-%d").unique()) == ["2024-05-01", "2024-05-02", "2026-01-10"]
    assert (lu[lu["report_date"] == "2026-01-10"]["kmt"].sum()) == pytest.approx(170)                                   # Southport day untouched
    prov = store.load_table("lineup_provenance", db)
    assert len(prov) == 2 and set(prov["provider"]) == {"APS"}
    assert ap.backfill(str(folder), db, since="2024-01-01")["to_load"] == 0                                              # idempotent


def test_a_region_not_reported_on_aps_dates_is_not_read_as_empty():
    rows = [snap(d, "Texas Gulf", 5, 150) for d in ("2026-01-02", "2026-01-12")] + [snap(d, "US Gulf", 20, 700) for d in ("2026-01-02", "2026-01-07", "2026-01-12")]
    l = pd.concat(rows)
    d_wrong, m_wrong = ch.lineup_detail(l, "vessels")
    unc = {("US", "Texas Gulf", pd.Timestamp("2026-01-07"))}
    d_ok, m_ok = ch.lineup_detail(l, "vessels", unc)
    tx = lambda d: ch.rollup(d[d["region"] == "Texas Gulf"], ["region"]).iloc[0]
    assert tx(d_ok)["now"] == 5 and tx(d_ok)["prev_wk"] == 5                      # compares with Jan 2, not with a false zero on Jan 7
    assert pd.Timestamp(m_ok[m_ok["region"] == "Texas Gulf"].iloc[0]["wk_date"]) == pd.Timestamp("2026-01-02")
